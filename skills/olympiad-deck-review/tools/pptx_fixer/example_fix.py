"""Example fix script. Copy it, rename it, and replace the fix list with yours.

Usage:
    python example_fix.py "My Deck.pptx" "My Deck (reviewed v1 - content fixes).pptx"

The output file name must be NEW. The script refuses to overwrite its input.
Every fix below uses made-up content: replace slide numbers and text with
what is really in YOUR deck, found with dump_deck.py and the rendered images.
"""
import sys
from pptx import Presentation
import pptx_fixer as fx

src, out = sys.argv[1], sys.argv[2]
if src == out:
    raise SystemExit('Pick a new output name. Never overwrite the original.')

prs = Presentation(src)
S = lambda n: prs.slides[n - 1]          # S(5) is slide 5

# 1. Change a word that is repeated on several slides (question + each solution step).
#    expected_each=1 means "stop if this is not found exactly once on each slide".
fx.replace_text_in_slides(prs, range(5, 9), 'a room temperature', 'room temperature',
                          expected_each=1, tag='Q3 stem: stray "a"')

# 2. Rewrite a whole text box, keeping its look. '' makes a blank line.
box = fx.find_shape(S(6), contains='So option A is already')
fx.set_lines(box, ['So far: in the first mixture, salt is the solute.', '',
                   'This rules out option D.'])

# 3. A green "correct" box sits one row too low: move it up to option B's row.
#    Find it by its shape type and approximate position (inches from top-left).
tick = fx.find_shape(S(10), geom='roundRect', at=(0.9, 3.4))
fx.move_shape(tick, top_in=2.9)

# 4. Add a second cross mark and make it appear with the first click,
#    so the presenter does not need an extra click.
cross = fx.find_shape(S(7), geom='mathMultiply', at=(4.9, 2.2))
second = fx.clone_shape(S(7), cross, top_in=3.4)
fx.animate_with_click(S(7), 0, second)   # 0 = appears together with click 1

# 5. Swap a picture for a corrected one (same position, size and animation).
#    Make the new image the same shape (aspect ratio) as the old one.
# pic = next(s for s in S(12).shapes if s.shape_type == 13)
# fx.replace_image(S(12), pic, 'corrected_figure.png')

prs.save(out)
fx.print_log()
print('Saved:', out)
print('Next: render to PDF with LibreOffice and look at every slide you changed.')
