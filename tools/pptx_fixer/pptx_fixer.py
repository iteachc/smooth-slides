"""pptx_fixer: small, safe helpers for fixing content in a .pptx with python-pptx.

Why this exists
---------------
Most fixes to a practice-test deck are tiny: change a word, move a tick box,
swap a picture. PowerPoint splits sentences into several "runs" (pieces of
text with their own formatting), so a plain find-and-replace often misses text
or wipes the formatting. These helpers handle that.

Everything here edits an in-memory Presentation. Your script decides where to
save, and should ALWAYS save to a NEW file name (never overwrite the original).

Slide numbers are 1-based everywhere, matching what you see in PowerPoint.
"""
import copy
from lxml import etree
from pptx.util import Emu

NS_P = 'http://schemas.openxmlformats.org/presentationml/2006/main'
NS_A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
NS_R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
P_ = '{%s}' % NS_P
A_ = '{%s}' % NS_A
R_ = '{%s}' % NS_R
EMU_PER_INCH = 914400

# Everything the helpers do is recorded here so you can print a change report.
LOG = []


# ---------------------------------------------------------------- finding things

def walk(shapes):
    """Yield every shape, including shapes inside groups."""
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:  # group
            yield from walk(sh.shapes)


def text_frames(slide):
    """Yield (shape, text_frame) for every text box and table cell on a slide."""
    for sh in walk(slide.shapes):
        if sh.has_text_frame:
            yield sh, sh.text_frame
        if getattr(sh, 'has_table', False) and sh.has_table:
            for row in sh.table.rows:
                for cell in row.cells:
                    yield sh, cell.text_frame


def find_shape(slide, contains=None, geom=None, at=None, tol=0.08):
    """Find ONE shape. Stops with a clear message if nothing matches.

    contains: text the shape must contain
    geom:     preset shape name, e.g. 'mathMultiply' (a cross), 'rect', 'ellipse'
    at:       (left, top) in inches; matches within `tol` inches
    """
    for sh in walk(slide.shapes):
        if contains and not (sh.has_text_frame and contains in sh.text_frame.text):
            continue
        if geom:
            g = sh._element.find('.//' + A_ + 'prstGeom')
            if g is None or g.get('prst') != geom:
                continue
        if at and not (abs(sh.left / EMU_PER_INCH - at[0]) < tol
                       and abs(sh.top / EMU_PER_INCH - at[1]) < tol):
            continue
        return sh
    raise SystemExit(f'Shape not found: contains={contains!r} geom={geom!r} at={at!r}')


# ---------------------------------------------------------------- text edits

def replace_text(slide, old, new, expected=None, tag=''):
    """Replace `old` with `new` anywhere on the slide, even when `old` is split
    across several runs. Formatting of the first run it touches is kept.

    expected: if given, stop with an error unless exactly this many matches were
              replaced. Use it: it catches typos in `old` before they cost you.
    Returns the number of replacements.
    """
    count = 0
    for _sh, tf in text_frames(slide):
        for para in tf.paragraphs:
            runs = para.runs
            full = ''.join(r.text for r in runs)
            while old in full:
                start = full.index(old)
                end = start + len(old)
                pos, placed = 0, False
                for r in runs:
                    rs, re_ = pos, pos + len(r.text)
                    pos = re_
                    if re_ <= start or rs >= end:
                        continue
                    a, b = max(start, rs) - rs, min(end, re_) - rs
                    if not placed:
                        r.text = r.text[:a] + new + r.text[b:]
                        placed = True
                    else:
                        r.text = r.text[:a] + r.text[b:]
                full = ''.join(r.text for r in runs)
                count += 1
                if new and old in new:  # avoid an endless loop
                    break
    LOG.append((tag or old[:40], count))
    if expected is not None and count != expected:
        raise SystemExit(f'replace_text({old!r}): found {count}, expected {expected}')
    return count


def replace_text_in_slides(prs, slide_numbers, old, new, expected_each=None, tag=''):
    """Same replacement on several slides, e.g. a question stem repeated on slides 5-9."""
    total = 0
    for n in slide_numbers:
        total += replace_text(prs.slides[n - 1], old, new, expected_each,
                              f'{tag or old[:30]} (slide {n})')
    return total


def set_lines(shape, lines):
    """Replace ALL the text in a shape with `lines` (one paragraph each),
    keeping the formatting of the shape's first run. Use '' for a blank line."""
    body = shape.text_frame._txBody
    paras = body.findall(A_ + 'p')
    template = copy.deepcopy(paras[0])
    for extra in template.findall(A_ + 'r')[1:] + template.findall(A_ + 'br'):
        template.remove(extra)
    for p in paras:
        body.remove(p)
    for line in lines:
        para = copy.deepcopy(template)
        runs = para.findall(A_ + 'r')
        if runs:
            runs[0].find(A_ + 't').text = line
        body.append(para)
    LOG.append((f'set_lines on "{shape.name}"', len(lines)))


# ---------------------------------------------------------------- shapes

def move_shape(shape, top_in=None, left_in=None):
    """Move a shape (inches from the top-left of the slide)."""
    if top_in is not None:
        shape.top = Emu(int(top_in * EMU_PER_INCH))
    if left_in is not None:
        shape.left = Emu(int(left_in * EMU_PER_INCH))


def _next_shape_id(slide):
    ids = [int(x.get('id')) for x in slide._element.iter()
           if x.tag.endswith('}cNvPr') and (x.get('id') or '').isdigit()]
    return max(ids) + 1


def clone_shape(slide, shape, top_in=None, left_in=None):
    """Copy a shape (same formatting) onto the same slide, optionally moved.
    The copy is placed right after the original in the stacking order."""
    el = copy.deepcopy(shape._element)
    new_id = _next_shape_id(slide)
    el.find('.//' + P_ + 'cNvPr').set('id', str(new_id))
    shape._element.addnext(el)
    new = [s for s in slide.shapes if s.shape_id == new_id][0]
    move_shape(new, top_in, left_in)
    LOG.append((f'clone of "{shape.name}"', 1))
    return new


# ---------------------------------------------------------------- animation

def click_steps(slide):
    """The list of click steps in the slide's main animation sequence."""
    timing = slide._element.find(P_ + 'timing')
    if timing is None:
        raise SystemExit('This slide has no animations.')
    seq = timing.find('.//' + P_ + 'seq')
    return seq.find(P_ + 'cTn').find(P_ + 'childTnLst').findall(P_ + 'par')


def is_animated(slide, shape):
    timing = slide._element.find(P_ + 'timing')
    return timing is not None and any(
        t.get('spid') == str(shape.shape_id) for t in timing.iter(P_ + 'spTgt'))


def animate_with_click(slide, click_index, shape):
    """Make `shape` APPEAR together with the existing animation on click number
    `click_index` (0 = first click). It adds to that click; it does not create a
    new click, so the presenter's click count stays the same.
    Does nothing if the shape is already animated."""
    if is_animated(slide, shape):
        return
    first = click_steps(slide)[click_index].find(P_ + 'cTn').find(P_ + 'childTnLst').find(P_ + 'par')
    effects = first.find(P_ + 'cTn').find(P_ + 'childTnLst')
    used = [int(c.get('id')) for c in slide._element.iter(P_ + 'cTn') if (c.get('id') or '').isdigit()]
    a, b = max(used + [100]) + 1, max(used + [100]) + 2
    effects.append(etree.fromstring(
        f'<p:par xmlns:p="{NS_P}"><p:cTn id="{a}" presetID="1" presetClass="entr" presetSubtype="0" '
        f'fill="hold" nodeType="withEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
        f'<p:set><p:cBhvr><p:cTn id="{b}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>'
        f'<p:tgtEl><p:spTgt spid="{shape.shape_id}"/></p:tgtEl>'
        f'<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr>'
        f'<p:to><p:strVal val="visible"/></p:to></p:set></p:childTnLst></p:cTn></p:par>'))
    LOG.append((f'animate "{shape.name}" with click {click_index + 1}', 1))


# ---------------------------------------------------------------- images

def replace_image(slide, picture, new_image_path):
    """Swap the picture inside an existing picture shape for a new image file.
    The shape keeps its position, size and animations, so only the pixels change.
    Tip: make the new image the same aspect ratio as the old one, or it will look
    stretched. Uses a fresh image part, so other slides sharing the old picture
    are not affected."""
    _image_part, new_rid = slide.part.get_or_add_image_part(new_image_path)
    blip = picture._element.find('.//' + A_ + 'blip')
    blip.set(R_ + 'embed', new_rid)
    LOG.append((f'replace_image "{picture.name}" <- {new_image_path}', 1))


def extract_images(prs, out_dir):
    """Save every picture in the deck as slide<N>_<shape id>.<ext> so you can look at them."""
    import os
    os.makedirs(out_dir, exist_ok=True)
    for n, slide in enumerate(prs.slides, 1):
        for sh in walk(slide.shapes):
            if sh.shape_type == 13:  # picture
                path = os.path.join(out_dir, f'slide{n}_{sh.shape_id}.{sh.image.ext}')
                with open(path, 'wb') as f:
                    f.write(sh.image.blob)


# ---------------------------------------------------------------- slide info

def hidden_slides(prs):
    """Slide numbers (1-based) that are hidden in the slide show."""
    return [n for n, s in enumerate(prs.slides, 1) if s._element.get('show') == '0']


def pdf_page_for_slide(prs, slide_number):
    """LibreOffice leaves hidden slides out of the PDF, so page numbers shift.
    Returns the PDF page for a slide, or None if the slide is hidden."""
    hidden = hidden_slides(prs)
    if slide_number in hidden:
        return None
    return slide_number - sum(1 for h in hidden if h < slide_number)


def print_log():
    print('--- change report ---')
    for tag, n in LOG:
        print(f'{n:>3}  {tag}')
