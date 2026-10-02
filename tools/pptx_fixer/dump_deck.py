"""Print everything in a deck as text: slide by slide, with hidden slides flagged.

Usage:  python dump_deck.py deck.pptx

WARNING: shapes are listed in layer (z-order) order, NOT in the order you see
them on screen. Option letters (A, B, C, D) can come out shuffled. Use this
dump to find text and shapes; read answer keys from the RENDERED slide images.
"""
import sys
from pptx import Presentation
from pptx_fixer import walk, hidden_slides

prs = Presentation(sys.argv[1])
print('SLIDES:', len(prs.slides), '| hidden:', hidden_slides(prs) or 'none')
for i, s in enumerate(prs.slides, 1):
    print('=' * 58)
    print('SLIDE', i, '[HIDDEN]' if s._element.get('show') == '0' else '')
    for sh in walk(s.shapes):
        pos = f'(left {sh.left // 9525}, top {sh.top // 9525})' if sh.left is not None else ''
        if sh.shape_type == 13:
            print('  [PICTURE]', pos, sh.width // 9525, 'x', sh.height // 9525)
        if getattr(sh, 'has_table', False) and sh.has_table:
            for r in sh.table.rows:
                print('  T|', ' | '.join(c.text.replace('\n', ' / ') for c in r.cells))
        if sh.has_text_frame and sh.text_frame.text.strip():
            print(' ', pos, sh.text_frame.text.strip().replace('\n', ' / '))
    if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip():
        print('  --NOTES--', s.notes_slide.notes_text_frame.text.strip()[:300])
